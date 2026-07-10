from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from django.db import transaction
from openpyxl import load_workbook

from .models import Student


EXPECTED_HEADERS = {
    "nombres": "first_name",
    "apellidos": "last_name",
    "cedula": "document_id",
    "sexo": "gender",
    "sigerd": "sigerd_id",
    "nacimiento": "birth_date",
    "curso": "course",
}


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    enrollments_created: int = 0
    courses_created: int = 0
    errors: list[str] = field(default_factory=list)


@dataclass
class ImportPreview:
    rows: list[dict] = field(default_factory=list)
    total: int = 0
    to_create: int = 0
    to_update: int = 0
    to_skip: int = 0
    errors: list[str] = field(default_factory=list)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def clean_document(value: Any) -> str:
    return clean_text(value)


def parse_birth_date(value: Any):
    if value in (None, ""):
        return None
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
    return None


def serialize_date(value):
    if isinstance(value, date):
        return value.isoformat()
    return ""


def deserialize_date(value):
    if not value:
        return None
    return datetime.strptime(value, "%Y-%m-%d").date()


def normalize_gender(value: Any) -> str:
    gender = clean_text(value).upper()
    if gender in {"M", "F"}:
        return gender
    if gender.startswith("MASC"):
        return "M"
    if gender.startswith("FEM"):
        return "F"
    return gender[:20]


def split_course_section(value: Any):
    raw_value = clean_text(value).upper().replace("-", " ")
    if not raw_value:
        return "", ""

    compact = raw_value.replace(" ", "")
    if len(compact) >= 2 and compact[-1].isalpha() and any(ch.isdigit() for ch in compact[:-1]):
        return compact[:-1], compact[-1]

    parts = raw_value.split()
    if len(parts) >= 2:
        return parts[0], parts[1]

    return raw_value, ""


def get_headers(ws):
    headers = {}
    for cell in ws[1]:
        header = clean_text(cell.value).lower()
        if header:
            headers[header] = cell.column
    missing = [header for header in EXPECTED_HEADERS if header not in headers]
    return headers, missing


def find_existing_student(data):
    document_id = data["document_id"]
    sigerd_id = data["sigerd_id"]
    birth_date = data["birth_date"]

    if document_id:
        student = Student.objects.filter(document_id__iexact=document_id).first()
        if student:
            return student

    if sigerd_id:
        student = Student.objects.filter(sigerd_id__iexact=sigerd_id).first()
        if student:
            return student

    if birth_date:
        return Student.objects.filter(
            first_name__iexact=data["first_name"],
            last_name__iexact=data["last_name"],
            birth_date=birth_date,
        ).first()

    return None


def row_key(data):
    if data["document_id"]:
        return f"cedula:{data['document_id'].lower()}"
    if data["sigerd_id"]:
        return f"sigerd:{data['sigerd_id'].lower()}"
    return f"name:{data['first_name'].lower()}|{data['last_name'].lower()}|{data['birth_date']}"


def read_student_rows_from_excel(file_obj, update_existing: bool = True) -> ImportPreview:
    preview = ImportPreview()
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    ws = workbook["Estudiantes"] if "Estudiantes" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    headers, missing = get_headers(ws)
    if missing:
        preview.errors.append("Faltan columnas requeridas: " + ", ".join(missing))
        return preview

    seen = set()
    for row_number in range(2, ws.max_row + 1):
        data = {}
        for header, field_name in EXPECTED_HEADERS.items():
            data[field_name] = ws.cell(row_number, headers[header]).value

        data = {
            "first_name": clean_text(data["first_name"]),
            "last_name": clean_text(data["last_name"]),
            "document_id": clean_document(data["document_id"]),
            "gender": normalize_gender(data["gender"]),
            "sigerd_id": clean_text(data["sigerd_id"]),
            "birth_date": parse_birth_date(data["birth_date"]),
            "course": clean_text(data["course"]).upper(),
        }
        data["course"], data["section"] = split_course_section(data["course"])

        if not data["first_name"] and not data["last_name"]:
            preview.to_skip += 1
            continue

        row = {
            "row_number": row_number,
            "first_name": data["first_name"],
            "last_name": data["last_name"],
            "document_id": data["document_id"],
            "gender": data["gender"],
            "sigerd_id": data["sigerd_id"],
            "birth_date": serialize_date(data["birth_date"]),
            "course": data["course"],
            "section": data["section"],
            "status": "",
            "message": "",
        }

        if not data["first_name"] or not data["last_name"]:
            row["status"] = "error"
            row["message"] = "Nombres y apellidos son obligatorios."
            preview.errors.append(f"Fila {row_number}: {row['message']}")
            preview.to_skip += 1
            preview.rows.append(row)
            continue

        key = row_key(data)
        if key in seen:
            row["status"] = "skip"
            row["message"] = "Registro duplicado dentro del archivo."
            preview.errors.append(f"Fila {row_number}: {row['message']}")
            preview.to_skip += 1
            preview.rows.append(row)
            continue
        seen.add(key)

        student = find_existing_student(data)
        if student:
            if update_existing:
                row["status"] = "update"
                row["message"] = f"Actualizara estudiante existente #{student.pk}."
                preview.to_update += 1
            else:
                row["status"] = "skip"
                row["message"] = f"Ya existe estudiante #{student.pk}."
                preview.to_skip += 1
        else:
            row["status"] = "create"
            row["message"] = "Creara estudiante nuevo."
            preview.to_create += 1

        preview.rows.append(row)

    preview.total = len(preview.rows)
    return preview


def row_to_data(row):
    return {
        "first_name": row["first_name"],
        "last_name": row["last_name"],
        "document_id": row["document_id"],
        "gender": row["gender"],
        "sigerd_id": row["sigerd_id"],
        "birth_date": deserialize_date(row["birth_date"]),
        "course": row["course"],
        "section": row["section"],
    }


@transaction.atomic
def commit_student_import(preview_rows: list[dict], school_year: str, update_existing: bool = True) -> ImportResult:
    result = ImportResult()

    for row in preview_rows:
        if row["status"] in {"error", "skip"}:
            result.skipped += 1
            continue

        data = row_to_data(row)
        student = find_existing_student(data)
        if student:
            if update_existing:
                student.first_name = data["first_name"]
                student.last_name = data["last_name"]
                student.document_id = data["document_id"]
                student.gender = data["gender"]
                student.sigerd_id = data["sigerd_id"]
                student.birth_date = data["birth_date"]
                student.new_admission = True
                student.active = True
                student.save()
                result.updated += 1
            else:
                result.skipped += 1
                continue
        else:
            student = Student.objects.create(
                first_name=data["first_name"],
                last_name=data["last_name"],
                document_id=data["document_id"],
                gender=data["gender"],
                sigerd_id=data["sigerd_id"],
                birth_date=data["birth_date"],
                new_admission=True,
                active=True,
            )
            result.created += 1

    return result


def import_students_from_excel(file_obj, school_year: str, update_existing: bool = True) -> ImportResult:
    preview = read_student_rows_from_excel(file_obj)
    if preview.errors:
        return ImportResult(skipped=preview.to_skip, errors=preview.errors)
    return commit_student_import(preview.rows, school_year, update_existing)
