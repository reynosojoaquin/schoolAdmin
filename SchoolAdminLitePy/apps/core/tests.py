import shutil
import tempfile
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Permission, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from .forms import AdministrativeEmployeeForm, StudentTransferForm
from .models import (
    AdministrativeEmployee,
    Attendance,
    City,
    Course,
    EmployeePosition,
    EmployeeType,
    Enrollment,
    Grade,
    Nationality,
    Province,
    Section,
    Student,
    Subject,
    SystemConfiguration,
    Teacher,
    TeachingAssignment,
)


TEST_MEDIA_ROOT = tempfile.mkdtemp()


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class SystemConfigurationTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEST_MEDIA_ROOT, ignore_errors=True)

    def test_login_displays_default_logo_and_session_label(self):
        response = self.client.get(reverse("core:login"))

        self.assertContains(response, "logo-centro-default")
        self.assertContains(response, "Iniciar sesi&oacute;n", html=True)

    def test_only_system_managers_can_update_branding(self):
        regular_user = User.objects.create_user("regular", password="test-password")
        self.client.force_login(regular_user)
        self.assertEqual(self.client.get(reverse("core:system_configuration")).status_code, 403)

        admin = User.objects.create_superuser("admin", "admin@example.com", "test-password")
        self.client.force_login(admin)
        response = self.client.post(
            reverse("core:system_configuration"),
            {
                "institution_name": "Centro Educativo de Prueba",
                "current_school_year": "2026-2027",
                "logo": SimpleUploadedFile("logo.png", b"fake-png-content", content_type="image/png"),
            },
        )

        self.assertRedirects(response, reverse("core:system_configuration"))
        configuration = SystemConfiguration.objects.get(pk=1)
        self.assertEqual(configuration.institution_name, "Centro Educativo de Prueba")
        self.assertTrue(configuration.logo.name.startswith("branding/logo"))

    def test_authenticated_layout_uses_configured_logo(self):
        admin = User.objects.create_superuser("admin", "admin@example.com", "test-password")
        SystemConfiguration.objects.create(
            logo=SimpleUploadedFile("custom-logo.jpg", b"fake-jpg-content", content_type="image/jpeg")
        )
        self.client.force_login(admin)

        response = self.client.get(reverse("core:dashboard"))

        self.assertContains(response, "/media/branding/custom-logo.jpg")
        self.assertContains(response, reverse("core:system_configuration"))


class AdministrativeEmployeeCatalogTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser("catalog-admin", "admin@example.com", "password")
        cls.nationality = Nationality.objects.create(name="Nacionalidad de prueba")
        cls.other_nationality = Nationality.objects.create(name="Otra nacionalidad de prueba")
        cls.province = Province.objects.create(name="Provincia de prueba", nationality=cls.nationality)
        cls.other_province = Province.objects.create(name="Otra provincia de prueba", nationality=cls.other_nationality)
        cls.city = City.objects.create(province=cls.province, name="Ciudad de prueba")
        cls.other_city = City.objects.create(province=cls.other_province, name="Otra ciudad")
        cls.employee_type = EmployeeType.objects.create(name="Tipo de prueba")
        cls.other_employee_type = EmployeeType.objects.create(name="Otro tipo de prueba")
        cls.position = EmployeePosition.objects.create(
            employee_type=cls.employee_type,
            name="Posición de prueba",
        )
        cls.other_position = EmployeePosition.objects.create(
            employee_type=cls.other_employee_type,
            name="Otra posición",
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_city_endpoint_filters_by_province(self):
        response = self.client.get(
            reverse("core:birth_cities_options"),
            {"province": self.province.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["cities"], [{"id": self.city.pk, "text": self.city.name}])

    def test_province_endpoint_filters_by_nationality(self):
        response = self.client.get(
            reverse("core:birth_provinces_options"),
            {"nationality": self.nationality.pk},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["provinces"], [{"id": self.province.pk, "text": self.province.name}])

    def test_position_endpoint_filters_by_employee_type(self):
        response = self.client.get(
            reverse("core:employee_positions_options"),
            {"employee_type": self.employee_type.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.json()["positions"],
            [{"id": self.position.pk, "text": self.position.name}],
        )

    def test_form_limits_dependent_catalogs(self):
        form = AdministrativeEmployeeForm(
            data={
                "nationality": self.nationality.pk,
                "birth_province": self.province.pk,
                "employee_type": self.employee_type.pk,
            }
        )

        self.assertQuerySetEqual(form.fields["birth_province"].queryset, [self.province])
        self.assertQuerySetEqual(form.fields["birthplace"].queryset, [self.city])
        self.assertQuerySetEqual(form.fields["position"].queryset, [self.position])

    def test_gender_is_a_select_with_expected_options(self):
        form = AdministrativeEmployeeForm()

        self.assertEqual(
            list(form.fields["gender"].choices),
            [("", "---------"), ("Masculino", "Masculino"), ("Femenino", "Femenino")],
        )
        self.assertEqual(form.fields["active"].label, "Personal administrativo habilitado")

    def test_inactive_employee_is_visible_only_to_administrator(self):
        inactive_employee = AdministrativeEmployee.objects.create(
            first_name="Empleado",
            last_name="Inactivo",
            active=False,
        )
        regular_user = User.objects.create_user("people-viewer", password="password")
        regular_user.user_permissions.add(
            Permission.objects.get(codename="manage_people", content_type__app_label="core")
        )

        self.client.force_login(self.admin)
        admin_response = self.client.get(reverse("core:employee_list"))
        self.assertContains(admin_response, "Inactivo, Empleado")

        self.client.force_login(regular_user)
        regular_response = self.client.get(reverse("core:employee_list"))
        self.assertNotContains(regular_response, str(inactive_employee))

    def test_employee_active_field_can_disable_employee(self):
        response = self.client.post(
            reverse("core:employee_create"),
            {
                "first_name": "Empleado",
                "last_name": "Deshabilitado",
                "active": "",
            },
        )

        self.assertEqual(response.status_code, 302)
        employee = AdministrativeEmployee.objects.get(first_name="Empleado")
        self.assertFalse(employee.active)

    def test_employee_list_assigns_order_numbers(self):
        first = AdministrativeEmployee.objects.create(first_name="Ana", last_name="Alba")
        second = AdministrativeEmployee.objects.create(first_name="Bea", last_name="Bravo")

        response = self.client.get(reverse("core:employee_list"))

        self.assertEqual(
            [(employee, employee.order_number) for employee in response.context["object_list"]],
            [(first, 1), (second, 2)],
        )

    def test_employee_phone_is_saved_and_can_be_updated(self):
        response = self.client.post(reverse("core:employee_create"), {
            "first_name": "Ana", "last_name": "Perez", "phone": "809-555-0101",
        })
        self.assertEqual(response.status_code, 302)
        employee = AdministrativeEmployee.objects.get(first_name="Ana")
        self.assertEqual(employee.phones.get().number, "809-555-0101")
        response = self.client.post(reverse("core:employee_update", args=[employee.pk]), {
            "first_name": "Ana", "last_name": "Perez", "phone": "809-555-0102",
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(employee.phones.get().number, "809-555-0102")

    def test_employee_birth_date_is_shown_and_saved_on_edit(self):
        employee = AdministrativeEmployee.objects.create(
            first_name="Ana", last_name="Perez", birth_date=date(1985, 4, 12)
        )
        url = reverse("core:employee_update", args=[employee.pk])
        response = self.client.get(url)
        self.assertContains(response, 'type="date"')
        self.assertContains(response, 'value="1985-04-12"')

        response = self.client.post(url, {
            "first_name": "Ana", "last_name": "Perez", "birth_date": "1986-05-13",
        })
        self.assertRedirects(response, reverse("core:employee_list"))
        employee.refresh_from_db()
        self.assertEqual(employee.birth_date, date(1986, 5, 13))


    def test_form_rejects_city_or_position_from_another_parent(self):
        form = AdministrativeEmployeeForm(
            data={
                "first_name": "Ana",
                "last_name": "Pérez",
                "nationality": self.nationality.pk,
                "birth_province": self.province.pk,
                "birthplace": self.other_city.pk,
                "employee_type": self.employee_type.pk,
                "position": self.other_position.pk,
            }
        )

        self.assertFalse(form.is_valid())
        self.assertIn("birthplace", form.errors)
        self.assertIn("position", form.errors)

    def test_form_rejects_province_from_another_nationality(self):
        form = AdministrativeEmployeeForm(data={
            "first_name": "Ana", "last_name": "Perez",
            "nationality": self.nationality.pk,
            "birth_province": self.other_province.pk,
        })
        self.assertFalse(form.is_valid())
        self.assertIn("birth_province", form.errors)


class SectionEnrollmentTests(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_superuser("enrollment-admin", "admin@example.com", "password"))
        self.course = Course.objects.create(name="Curso pruebas inscripcion")
        self.section = Section.objects.create(course=self.course, name="A", school_year="2026-2027")
        configuration = SystemConfiguration.get_solo()
        configuration.current_school_year = "2026-2027"
        configuration.save(update_fields=["current_school_year"])
        self.student = Student.objects.create(first_name="Eva", last_name="Perez", new_admission=True)

    def test_enrollment_is_assigned_to_selected_section(self):
        response = self.client.post(reverse("core:section_enrollment", args=[self.section.pk]), {
            "students": [self.student.pk],
        })
        self.assertRedirects(response, reverse("core:section_students", args=[self.section.pk]))
        enrollment = Enrollment.objects.get(student=self.student)
        self.assertEqual(enrollment.section, self.section)
        self.student.refresh_from_db()
        self.assertFalse(self.student.new_admission)

    def test_section_delete_identifies_protected_relationship(self):
        Subject.objects.create(section=self.section, name="Asignatura protegida")

        response = self.client.post(
            reverse("core:section_delete", args=[self.section.pk]),
            follow=True,
        )

        self.assertRedirects(response, reverse("core:section_list"))
        self.assertTrue(Section.objects.filter(pk=self.section.pk).exists())
        self.assertContains(response, "asignatura")


class StudentSectionOrderTests(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_superuser("order-admin", "admin@example.com", "password"))
        configuration = SystemConfiguration.get_solo()
        configuration.current_school_year = "2026-2027"
        configuration.save(update_fields=["current_school_year"])
        course = Course.objects.create(name="Curso orden pruebas")
        self.section_a = Section.objects.create(course=course, name="A", school_year="2026-2027")
        self.section_b = Section.objects.create(course=course, name="B", school_year="2026-2027")
        self.zeta = Student.objects.create(first_name="Zoe", last_name="Zapata Cruz")
        self.alba = Student.objects.create(first_name="Ana", last_name="Alba Diaz")
        self.brito = Student.objects.create(first_name="Bea", last_name="Brito Santos")
        for student, section in ((self.zeta, self.section_a), (self.alba, self.section_a), (self.brito, self.section_b)):
            Enrollment.objects.create(student=student, course=course, section=section, school_year="2026-2027")

    def test_section_numbers_follow_first_surname_and_restart_per_section(self):
        response = self.client.get(reverse("core:section_students", args=[self.section_a.pk]))
        self.assertEqual([(item.student, item.order_number) for item in response.context["object_list"]], [
            (self.alba, 1), (self.zeta, 2),
        ])
        response = self.client.get(reverse("core:section_students", args=[self.section_b.pk]))
        self.assertEqual(response.context["object_list"][0].order_number, 1)

    def test_search_keeps_section_number_in_student_index(self):
        response = self.client.get(reverse("core:student_list"), {"q": "Zapata"})
        self.assertEqual(response.context["object_list"][0].order_number, 2)
        self.assertEqual(response.context["object_list"][0].current_section, self.section_a)

    def test_course_index_groups_section_numbers(self):
        response = self.client.get(reverse("core:course_students", args=[self.section_a.course_id]))
        self.assertEqual(
            [(item.section.name, item.order_number) for item in response.context["object_list"]],
            [("A", 1), ("A", 2), ("B", 1)],
        )

    def test_attendance_only_lists_selected_section_with_order_numbers(self):
        response = self.client.get(reverse("core:section_attendance", args=[self.section_a.pk]))

        self.assertEqual(
            [(row["enrollment"].student, row["order_number"]) for row in response.context["rows"]],
            [(self.alba, 1), (self.zeta, 2)],
        )
        self.assertNotContains(response, self.brito.first_name)
        self.assertContains(response, "No. 1")

    def test_attendance_post_does_not_create_records_for_another_section(self):
        enrollment_a = Enrollment.objects.get(student=self.alba)
        enrollment_b = Enrollment.objects.get(student=self.brito)

        self.client.post(
            reverse("core:section_attendance", args=[self.section_a.pk]),
            {
                "date": "2026-10-01",
                f"status_{enrollment_a.pk}": Attendance.ABSENT,
                f"status_{enrollment_b.pk}": Attendance.LATE,
            },
        )

        self.assertTrue(Attendance.objects.filter(enrollment=enrollment_a, date="2026-10-01").exists())
        self.assertFalse(Attendance.objects.filter(enrollment=enrollment_b, date="2026-10-01").exists())

    def test_gradebook_only_lists_students_from_assignment_section(self):
        teacher = Teacher.objects.create(first_name="Pablo", last_name="Martinez")
        subject = Subject.objects.create(section=self.section_a, name="Matematicas")
        assignment = TeachingAssignment.objects.create(
            section=self.section_a,
            subject=subject,
            teacher=teacher,
        )

        response = self.client.get(reverse("core:gradebook", args=[assignment.pk]))

        self.assertEqual(
            [row["enrollment"].student for row in response.context["rows"]],
            [self.alba, self.zeta],
        )
        self.assertNotContains(response, self.brito.first_name)


class GradeRecoveryTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("recovery-admin", "recovery@example.com", "password")
        self.client.force_login(self.admin)
        configuration = SystemConfiguration.get_solo()
        configuration.current_school_year = "2026-2027"
        configuration.save(update_fields=["current_school_year"])
        self.course = Course.objects.create(name="Curso recuperacion")
        self.section = Section.objects.create(course=self.course, name="A", school_year="2026-2027")
        self.student = Student.objects.create(first_name="Ana", last_name="Alba")
        self.enrollment = Enrollment.objects.create(
            student=self.student,
            course=self.course,
            section=self.section,
            school_year="2026-2027",
        )
        self.teacher = Teacher.objects.create(first_name="Docente", last_name="Prueba")
        self.subject = Subject.objects.create(section=self.section, name="Matematicas")
        self.assignment = TeachingAssignment.objects.create(
            section=self.section,
            subject=self.subject,
            teacher=self.teacher,
        )

    def test_recovery_replaces_only_the_lowest_grade_when_pair_average_is_below_70(self):
        grade = Grade.objects.create(
            enrollment=self.enrollment,
            subject=self.subject,
            period_1=50,
            period_2=80,
            recovery_1=75,
            period_3=65,
            period_4=60,
            recovery_2=55,
        )

        self.assertEqual(grade.effective_periods, [Decimal("75"), Decimal("80"), Decimal("65"), Decimal("60")])
        self.assertEqual(grade.average, Decimal("70"))

    def test_recovery_is_not_applied_when_pair_average_is_at_least_70(self):
        grade = Grade.objects.create(
            enrollment=self.enrollment,
            subject=self.subject,
            period_1=60,
            period_2=80,
            recovery_1=95,
        )

        self.assertEqual(grade.effective_periods[:2], [Decimal("60"), Decimal("80")])

    def test_template_contains_recovery_columns_and_imports_them(self):
        response = self.client.get(reverse("core:grade_template", args=[self.assignment.pk]))
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook["Calificaciones"]
        headers = {cell.value: cell.column for cell in sheet[11] if cell.value}

        self.assertIn("c1_r1", headers)
        self.assertIn("c1_r2", headers)
        self.assertIn("<70", sheet.cell(row=12, column=29).value)

        sheet.cell(row=12, column=headers["c1_p1"]).value = 50
        sheet.cell(row=12, column=headers["c1_p2"]).value = 80
        sheet.cell(row=12, column=headers["c1_r1"]).value = 75
        output = BytesIO()
        workbook.save(output)
        upload = SimpleUploadedFile(
            "calificaciones.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        preview_response = self.client.post(
            reverse("core:grade_import"),
            {
                "teaching_assignment": self.assignment.pk,
                "excel_file": upload,
                "action": "preview",
            },
        )
        self.assertEqual(preview_response.status_code, 200)
        self.client.post(
            reverse("core:grade_import"),
            {"teaching_assignment": self.assignment.pk, "action": "confirm"},
        )

        grade = Grade.objects.get(
            enrollment=self.enrollment,
            subject=self.subject,
            subject_competency__isnull=False,
            recovery_1=75,
        )
        self.assertEqual(grade.recovery_1, Decimal("75"))


class TeacherListOrderTests(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_superuser("teacher-order-admin", "teacher-admin@example.com", "password"))

    def test_teacher_list_numbers_follow_name_order(self):
        alpha = Teacher.objects.create(first_name="Ordenado", last_name="Alfa")
        bravo = Teacher.objects.create(first_name="Ordenado", last_name="Bravo")
        charlie = Teacher.objects.create(first_name="Ordenado", last_name="Charlie")

        response = self.client.get(reverse("core:teacher_list"), {"q": "Ordenado"})

        self.assertEqual(
            [(teacher, teacher.order_number) for teacher in response.context["object_list"]],
            [(alpha, 1), (bravo, 2), (charlie, 3)],
        )


class UserCrudTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("user-admin", "user-admin@example.com", "password")
        self.client.force_login(self.admin)

    def test_user_list_assigns_order_numbers(self):
        User.objects.create_user("alpha-user", password="password123")
        User.objects.create_user("bravo-user", password="password123")

        response = self.client.get(reverse("core:user_access_list"), {"q": "-user"})

        self.assertEqual(
            [(user.username, user.order_number) for user in response.context["object_list"]],
            [("alpha-user", 1), ("bravo-user", 2)],
        )

    def test_user_can_be_created_with_password(self):
        response = self.client.post(reverse("core:user_create"), {
            "username": "created-user",
            "first_name": "Usuario",
            "last_name": "Creado",
            "email": "created@example.com",
            "password1": "strong-pass-123",
            "password2": "strong-pass-123",
        })

        self.assertRedirects(response, reverse("core:user_access_list"))
        created_user = User.objects.get(username="created-user")
        self.assertTrue(created_user.check_password("strong-pass-123"))

    def test_user_can_be_deleted_but_not_self(self):
        user = User.objects.create_user("deletable-user", password="password123")

        response = self.client.post(reverse("core:user_delete", args=[user.pk]))

        self.assertRedirects(response, reverse("core:user_access_list"))
        self.assertFalse(User.objects.filter(pk=user.pk).exists())
        self.assertEqual(self.client.post(reverse("core:user_delete", args=[self.admin.pk])).status_code, 404)


class InactivePeopleVisibilityTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("people-admin", "people-admin@example.com", "password")
        self.viewer = User.objects.create_user("people-viewer", password="password")
        permission = Permission.objects.get(codename="view_all_people")
        self.viewer.user_permissions.add(permission)
        self.active_teacher = Teacher.objects.create(first_name="Activo", last_name="Docente", active=True)
        self.inactive_teacher = Teacher.objects.create(first_name="Inactivo", last_name="Docente", active=False)
        self.active_student = Student.objects.create(first_name="Activo", last_name="Estudiante", active=True)
        self.inactive_student = Student.objects.create(first_name="Inactivo", last_name="Estudiante", active=False)

    def test_non_admin_users_do_not_see_inactive_people(self):
        self.client.force_login(self.viewer)
        teacher_response = self.client.get(reverse("core:teacher_list"))
        student_response = self.client.get(reverse("core:student_list"))

        self.assertIn(self.active_teacher, teacher_response.context["object_list"])
        self.assertNotIn(self.inactive_teacher, teacher_response.context["object_list"])
        self.assertIn(self.active_student, student_response.context["object_list"])
        self.assertNotIn(self.inactive_student, student_response.context["object_list"])

    def test_administrators_can_see_inactive_people(self):
        self.client.force_login(self.admin)

        self.assertIn(self.inactive_teacher, self.client.get(reverse("core:teacher_list")).context["object_list"])
        self.assertIn(self.inactive_student, self.client.get(reverse("core:student_list")).context["object_list"])


class StudentTransferTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("transfer-admin", "admin@example.com", "password")
        self.client.force_login(self.admin)
        self.school_year = "2026-2027"
        configuration = SystemConfiguration.get_solo()
        configuration.current_school_year = self.school_year
        configuration.save(update_fields=["current_school_year"])

        self.course_1 = Course.objects.create(name="1ro")
        self.section_1a = Section.objects.create(course=self.course_1, name="A", school_year=self.school_year)
        self.section_1b = Section.objects.create(course=self.course_1, name="B", school_year=self.school_year)

        self.course_2 = Course.objects.create(name="2do")
        self.section_2a = Section.objects.create(course=self.course_2, name="A", school_year=self.school_year)

        self.student = Student.objects.create(first_name="Carlos", last_name="Santana")
        self.enrollment = Enrollment.objects.create(
            student=self.student,
            course=self.course_1,
            section=self.section_1a,
            school_year=self.school_year,
            active=True,
        )

    def test_form_binds_and_validates_section_correctly(self):
        form = StudentTransferForm(
            data={"course": self.course_1.pk, "section": self.section_1b.pk},
            school_year=self.school_year,
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["course"], self.course_1)
        self.assertEqual(form.cleaned_data["section"], self.section_1b)

    def test_transfer_between_sections_in_same_course(self):
        subject_a = Subject.objects.create(section=self.section_1a, name="Matematicas")
        subject_b = Subject.objects.create(section=self.section_1b, name="Matematicas")
        grade = Grade.objects.create(enrollment=self.enrollment, subject=subject_a, period_1=90)

        url = reverse("core:student_transfer", args=[self.student.pk])
        next_url = reverse("core:section_students", args=[self.section_1a.pk])
        response = self.client.post(f"{url}?next={next_url}", {
            "course": self.course_1.pk,
            "section": self.section_1b.pk,
            "next": next_url,
        })
        self.assertRedirects(response, next_url)

        self.enrollment.refresh_from_db()
        self.assertTrue(self.enrollment.active)
        self.assertEqual(self.enrollment.section, self.section_1b)
        self.assertEqual(self.enrollment.course, self.course_1)

        grade.refresh_from_db()
        self.assertEqual(grade.subject, subject_b)

    def test_transfer_to_different_course(self):
        url = reverse("core:student_transfer", args=[self.student.pk])
        response = self.client.post(url, {
            "course": self.course_2.pk,
            "section": self.section_2a.pk,
        })
        self.assertRedirects(response, reverse("core:student_list"))

        self.enrollment.refresh_from_db()
        self.assertFalse(self.enrollment.active)

        new_enrollment = Enrollment.objects.get(
            student=self.student,
            course=self.course_2,
            school_year=self.school_year,
        )
        self.assertTrue(new_enrollment.active)
        self.assertEqual(new_enrollment.section, self.section_2a)

    def test_student_without_active_enrollment_can_choose_course_and_section(self):
        student = Student.objects.create(first_name="Nuevo", last_name="Ingreso")
        url = reverse("core:student_transfer", args=[student.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Sin inscripcion activa")

        response = self.client.post(url, {
            "course": self.course_2.pk,
            "section": self.section_2a.pk,
            "school_year": self.school_year,
        })
        self.assertRedirects(response, reverse("core:student_list"))
        enrollment = Enrollment.objects.get(student=student, school_year=self.school_year)
        self.assertTrue(enrollment.active)
        self.assertEqual(enrollment.section, self.section_2a)

    def test_transfer_same_section_shows_warning(self):
        url = reverse("core:student_transfer", args=[self.student.pk])
        response = self.client.post(url, {
            "course": self.course_1.pk,
            "section": self.section_1a.pk,
        })
        self.assertRedirects(response, reverse("core:student_list"))
        self.enrollment.refresh_from_db()
        self.assertTrue(self.enrollment.active)
        self.assertEqual(self.enrollment.section, self.section_1a)
